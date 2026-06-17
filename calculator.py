import tkinter as tk
from tkinter import ttk, messagebox
import datetime
import os
import re
import json
import urllib.request
import urllib.parse
import threading

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

EXCEL_FILE = "せどり_利益計算.xlsx"
CONFIG_FILE = "keepa_config.json"

PLATFORM_DEFAULTS = {
    "Amazon FBA": {
        "referral_rate": 15.0,
        "fba_fee": 400,
        "shipping": 0,
    },
    "Amazon 自己発送": {
        "referral_rate": 15.0,
        "fba_fee": 0,
        "shipping": 600,
    },
    "メルカリ": {
        "referral_rate": 10.0,
        "fba_fee": 0,
        "shipping": 500,
    },
}


def load_api_key():
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, encoding="utf-8") as f:
            return json.load(f).get("api_key", "")
    return ""


def save_api_key(key):
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump({"api_key": key}, f, ensure_ascii=False)


def calc(purchase, selling, referral_rate, fba_fee, shipping):
    referral_fee = selling * (referral_rate / 100)
    profit = selling - purchase - referral_fee - fba_fee - shipping
    profit_rate = (profit / selling * 100) if selling > 0 else 0
    return referral_fee, profit, profit_rate


def calc_min_selling_price(purchase, referral_rate, fba_fee, shipping, min_profit_rate):
    """最低利益率を確保できる最低販売価格を計算する"""
    # profit = selling - purchase - selling*(rate/100) - fba - ship >= selling*(min_profit_rate/100)
    # selling*(1 - rate/100 - min_profit_rate/100) >= purchase + fba + ship
    denominator = 1 - (referral_rate / 100) - (min_profit_rate / 100)
    if denominator <= 0:
        return None
    return (purchase + fba_fee + shipping) / denominator


def extract_asin(text):
    text = text.strip()
    m = re.search(r"/(?:dp|gp/product)/([A-Z0-9]{10})", text)
    if m:
        return m.group(1)
    if re.fullmatch(r"[A-Z0-9]{10}", text):
        return text
    return None


def fetch_keepa_price(api_key, asin):
    """Keepa APIで現在のAmazon Japan最安値（マーケットプレイス新品）を取得"""
    url = (
        "https://api.keepa.com/product"
        f"?key={urllib.parse.quote(api_key)}"
        f"&domain=5&asin={asin}&history=1&stats=1"
    )
    req = urllib.request.Request(url, headers={"User-Agent": "SedoriTool/1.0"})
    with urllib.request.urlopen(req, timeout=10) as resp:
        data = json.loads(resp.read().decode())

    products = data.get("products")
    if not products:
        return None, None

    p = products[0]
    title = p.get("title", "")

    stats = p.get("stats", {})
    current = stats.get("current", [])

    # index 1 = マーケットプレイス新品最安値, 0 = Amazon本体価格
    price = None
    for idx in (1, 0):
        if idx < len(current) and current[idx] and current[idx] > 0:
            price = current[idx]
            break

    if price is not None:
        price = price / 100

    return price, title


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("せどり 利益計算ツール（最安値監視機能付き）")
        self.root.geometry("560x820")
        self.root.resizable(False, True)
        self.root.configure(bg="#F4F6F8")

        self.api_key = load_api_key()

        self.platform = tk.StringVar(value="Amazon FBA")
        self.asin_input = tk.StringVar()
        self.purchase = tk.StringVar()
        self.selling = tk.StringVar()
        self.referral_rate = tk.StringVar(value="15.0")
        self.fba_fee = tk.StringVar(value="400")
        self.shipping = tk.StringVar(value="0")
        self.product_title = tk.StringVar(value="")

        self._build()

        for v in (self.purchase, self.selling, self.referral_rate,
                  self.fba_fee, self.shipping):
            v.trace_add("write", self._refresh)

    # ── UI 構築 ──────────────────────────────────────────────────────────────

    def _build(self):
        tk.Label(self.root, text="せどり 利益計算ツール",
                 font=("Helvetica", 17, "bold"),
                 bg="#F4F6F8", fg="#2C3E50").pack(pady=(14, 2))

        nb = ttk.Notebook(self.root)
        nb.pack(fill="both", expand=True, padx=10, pady=6)

        tab1 = tk.Frame(nb, bg="#F4F6F8")
        tab2 = tk.Frame(nb, bg="#F4F6F8")
        nb.add(tab1, text="  利益計算  ")
        nb.add(tab2, text="  最安値監視  ")

        self._build_calc_tab(tab1)
        self._build_monitor_tab(tab2)

        if not EXCEL_AVAILABLE:
            tk.Label(self.root,
                     text="※ openpyxl 未インストール — pip install openpyxl で Excel保存が使えます",
                     font=("Helvetica", 8), bg="#F4F6F8", fg="#E74C3C").pack(pady=2)

    # ── 利益計算タブ ─────────────────────────────────────────────────────────

    def _build_calc_tab(self, parent):
        self._section_keepa(parent)
        self._section_platform(parent)
        self._section_input(parent)
        self._section_result(parent)
        self._section_buttons(parent)

    def _section_keepa(self, parent):
        frm = tk.LabelFrame(parent, text="Keepa 価格自動取得",
                            font=("Helvetica", 10), bg="#F4F6F8", padx=12, pady=8)
        frm.pack(fill="x", padx=20, pady=(6, 4))
        frm.columnconfigure(1, weight=1)

        tk.Label(frm, text="ASIN / URL", font=("Helvetica", 11),
                 bg="#F4F6F8", anchor="w").grid(row=0, column=0, sticky="w", pady=4)

        entry_frm = tk.Frame(frm, bg="#F4F6F8")
        entry_frm.grid(row=0, column=1, sticky="ew", padx=(6, 0))
        entry_frm.columnconfigure(0, weight=1)

        tk.Entry(entry_frm, textvariable=self.asin_input,
                 font=("Helvetica", 11), width=28).grid(row=0, column=0, sticky="ew")

        self.fetch_btn = tk.Button(entry_frm, text="価格取得",
                                   command=self._fetch_price,
                                   font=("Helvetica", 10, "bold"),
                                   bg="#2980B9", fg="white", relief="flat",
                                   padx=8, pady=2, cursor="hand2")
        self.fetch_btn.grid(row=0, column=1, padx=(6, 0))

        self.title_lbl = tk.Label(frm, textvariable=self.product_title,
                                  font=("Helvetica", 9), bg="#F4F6F8",
                                  fg="#555", wraplength=430, justify="left", anchor="w")
        self.title_lbl.grid(row=1, column=0, columnspan=2, sticky="w", pady=(2, 0))

        self.keepa_status = tk.StringVar(value="")
        tk.Label(frm, textvariable=self.keepa_status, font=("Helvetica", 9),
                 bg="#F4F6F8", fg="#888").grid(row=2, column=0, columnspan=2,
                                               sticky="w", pady=(0, 2))

    def _section_platform(self, parent):
        frm = tk.LabelFrame(parent, text="プラットフォーム",
                            font=("Helvetica", 10), bg="#F4F6F8", padx=12, pady=6)
        frm.pack(fill="x", padx=20, pady=4)
        for p in PLATFORM_DEFAULTS:
            tk.Radiobutton(frm, text=p, variable=self.platform, value=p,
                           command=self._on_platform, font=("Helvetica", 11),
                           bg="#F4F6F8").pack(side="left", padx=8)

    def _section_input(self, parent):
        frm = tk.LabelFrame(parent, text="入力項目",
                            font=("Helvetica", 10), bg="#F4F6F8", padx=16, pady=8)
        frm.pack(fill="x", padx=20, pady=4)
        frm.columnconfigure(1, weight=1)

        rows = [
            ("仕入れ値 (円)",   self.purchase,      False),
            ("販売価格 (円)",   self.selling,       False),
            ("手数料率 (%)",    self.referral_rate, True),
            ("FBA配送料 (円)",  self.fba_fee,       True),
            ("送料 (円)",       self.shipping,      True),
        ]
        for i, (label, var, muted) in enumerate(rows):
            fg = "#888" if muted else "#2C3E50"
            tk.Label(frm, text=label, font=("Helvetica", 11), bg="#F4F6F8",
                     fg=fg, anchor="w", width=18).grid(row=i, column=0,
                                                        sticky="w", pady=4)
            tk.Entry(frm, textvariable=var, font=("Helvetica", 12),
                     justify="right", width=14,
                     bg="#FFFFFF" if not muted else "#F0F0F0").grid(
                         row=i, column=1, sticky="e", pady=4, padx=(0, 4))

    def _section_result(self, parent):
        frm = tk.LabelFrame(parent, text="計算結果",
                            font=("Helvetica", 10), bg="#F4F6F8", padx=16, pady=8)
        frm.pack(fill="x", padx=20, pady=4)
        frm.columnconfigure(1, weight=1)

        self.r_referral = tk.StringVar(value="¥ 0")
        self.r_fba      = tk.StringVar(value="¥ 0")
        self.r_shipping = tk.StringVar(value="¥ 0")
        self.r_profit   = tk.StringVar(value="¥ 0")
        self.r_rate     = tk.StringVar(value="0.0 %")
        self.r_status   = tk.StringVar(value="")

        for i, (lbl, var) in enumerate([
            ("カテゴリー手数料", self.r_referral),
            ("FBA配送料",        self.r_fba),
            ("送料",             self.r_shipping),
        ]):
            tk.Label(frm, text=lbl, font=("Helvetica", 10), bg="#F4F6F8",
                     fg="#666", anchor="w").grid(row=i, column=0, sticky="w", pady=2)
            tk.Label(frm, textvariable=var, font=("Helvetica", 10),
                     bg="#F4F6F8", fg="#555").grid(row=i, column=1, sticky="e", pady=2)

        ttk.Separator(frm, orient="horizontal").grid(
            row=3, column=0, columnspan=2, sticky="ew", pady=6)

        tk.Label(frm, text="利益", font=("Helvetica", 13, "bold"),
                 bg="#F4F6F8").grid(row=4, column=0, sticky="w", pady=3)
        self.profit_lbl = tk.Label(frm, textvariable=self.r_profit,
                                   font=("Helvetica", 18, "bold"),
                                   bg="#F4F6F8", fg="#27AE60")
        self.profit_lbl.grid(row=4, column=1, sticky="e", pady=3)

        tk.Label(frm, text="利益率", font=("Helvetica", 13, "bold"),
                 bg="#F4F6F8").grid(row=5, column=0, sticky="w", pady=3)
        tk.Label(frm, textvariable=self.r_rate,
                 font=("Helvetica", 15, "bold"),
                 bg="#F4F6F8", fg="#2980B9").grid(row=5, column=1, sticky="e")

        self.status_lbl = tk.Label(frm, textvariable=self.r_status,
                                   font=("Helvetica", 16, "bold"), bg="#F4F6F8")
        self.status_lbl.grid(row=6, column=0, columnspan=2, pady=(6, 2))

    def _section_buttons(self, parent):
        frm = tk.Frame(parent, bg="#F4F6F8")
        frm.pack(pady=12)

        tk.Button(frm, text="  Excelに保存  ", command=self._save,
                  font=("Helvetica", 12, "bold"),
                  bg="#27AE60", fg="white", relief="flat",
                  padx=16, pady=8, cursor="hand2").pack(side="left", padx=8)

        tk.Button(frm, text="  クリア  ", command=self._clear,
                  font=("Helvetica", 12),
                  bg="#95A5A6", fg="white", relief="flat",
                  padx=16, pady=8, cursor="hand2").pack(side="left", padx=8)

    # ── 最安値監視タブ ────────────────────────────────────────────────────────

    def _build_monitor_tab(self, parent):
        # APIキー設定
        api_frm = tk.LabelFrame(parent, text="Keepa APIキー設定",
                                font=("Helvetica", 10), bg="#F4F6F8", padx=12, pady=8)
        api_frm.pack(fill="x", padx=16, pady=(8, 4))
        api_frm.columnconfigure(1, weight=1)

        tk.Label(api_frm, text="APIキー", font=("Helvetica", 10),
                 bg="#F4F6F8").grid(row=0, column=0, sticky="w")
        self.mon_apikey_var = tk.StringVar(value=self.api_key)
        tk.Entry(api_frm, textvariable=self.mon_apikey_var,
                 font=("Helvetica", 10), show="*", width=30).grid(
                     row=0, column=1, sticky="ew", padx=(6, 0))
        tk.Button(api_frm, text="保存", command=self._save_api_key,
                  font=("Helvetica", 9), bg="#7F8C8D", fg="white",
                  relief="flat", padx=6, pady=2).grid(row=0, column=2, padx=(4, 0))

        # 商品追加フォーム
        add_frm = tk.LabelFrame(parent, text="監視商品を追加",
                                font=("Helvetica", 10), bg="#F4F6F8", padx=12, pady=8)
        add_frm.pack(fill="x", padx=16, pady=4)
        add_frm.columnconfigure(1, weight=1)

        fields = [
            ("ASIN / URL",     "mon_asin"),
            ("仕入れ値 (円)",   "mon_purchase"),
            ("手数料率 (%)",    "mon_ref_rate"),
            ("FBA配送料 (円)", "mon_fba"),
            ("最低利益率 (%)", "mon_min_profit"),
        ]
        defaults = {"mon_ref_rate": "15.0", "mon_fba": "400", "mon_min_profit": "10.0"}
        for i, (lbl, attr) in enumerate(fields):
            tk.Label(add_frm, text=lbl, font=("Helvetica", 10),
                     bg="#F4F6F8", anchor="w", width=16).grid(
                         row=i, column=0, sticky="w", pady=3)
            var = tk.StringVar(value=defaults.get(attr, ""))
            setattr(self, attr + "_var", var)
            tk.Entry(add_frm, textvariable=var, font=("Helvetica", 10),
                     width=20, justify="right").grid(
                         row=i, column=1, sticky="e", pady=3, padx=(0, 4))

        tk.Button(add_frm, text="  リストに追加  ", command=self._add_monitor_item,
                  font=("Helvetica", 10, "bold"),
                  bg="#8E44AD", fg="white", relief="flat",
                  padx=10, pady=4, cursor="hand2").grid(
                      row=len(fields), column=0, columnspan=2, pady=(8, 2))

        # 監視リスト
        list_frm = tk.LabelFrame(parent, text="監視リスト",
                                 font=("Helvetica", 10), bg="#F4F6F8", padx=8, pady=6)
        list_frm.pack(fill="both", expand=True, padx=16, pady=4)

        cols = ("ASIN", "仕入値", "現在最安値", "推奨価格", "利益率", "判定")
        self.mon_tree = ttk.Treeview(list_frm, columns=cols, show="headings", height=6)
        col_widths = [100, 70, 90, 90, 70, 80]
        for col, w in zip(cols, col_widths):
            self.mon_tree.heading(col, text=col)
            self.mon_tree.column(col, width=w, anchor="center")

        vsb = ttk.Scrollbar(list_frm, orient="vertical", command=self.mon_tree.yview)
        self.mon_tree.configure(yscrollcommand=vsb.set)
        self.mon_tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        # タグで色付け
        self.mon_tree.tag_configure("ok",      background="#D5F5E3", foreground="#1E8449")
        self.mon_tree.tag_configure("warning", background="#FDEBD0", foreground="#A04000")
        self.mon_tree.tag_configure("ng",      background="#FADBD8", foreground="#C0392B")
        self.mon_tree.tag_configure("loading", foreground="#888888")

        # ボタン行
        btn_frm = tk.Frame(parent, bg="#F4F6F8")
        btn_frm.pack(pady=6)

        self.check_all_btn = tk.Button(
            btn_frm, text="  全件 最安値チェック  ",
            command=self._check_all,
            font=("Helvetica", 11, "bold"),
            bg="#E67E22", fg="white", relief="flat",
            padx=14, pady=6, cursor="hand2")
        self.check_all_btn.pack(side="left", padx=6)

        tk.Button(btn_frm, text="選択削除", command=self._delete_selected,
                  font=("Helvetica", 10),
                  bg="#95A5A6", fg="white", relief="flat",
                  padx=10, pady=6, cursor="hand2").pack(side="left", padx=6)

        self.mon_status_var = tk.StringVar(value="")
        tk.Label(parent, textvariable=self.mon_status_var,
                 font=("Helvetica", 9), bg="#F4F6F8", fg="#555").pack(pady=(0, 4))

        # 監視データ保持: {iid: {asin, purchase, ref_rate, fba, min_profit}}
        self._monitor_items = {}

    def _save_api_key(self):
        key = self.mon_apikey_var.get().strip()
        self.api_key = key
        save_api_key(key)
        self.mon_status_var.set("APIキーを保存しました")

    def _add_monitor_item(self):
        asin_raw = self.mon_asin_var.get().strip()
        asin = extract_asin(asin_raw)
        if not asin:
            messagebox.showwarning("入力エラー",
                "ASINまたはAmazonのURLを入力してください。\n例: B08XYZ1234")
            return
        try:
            purchase   = float(self.mon_purchase_var.get())
            ref_rate   = float(self.mon_ref_rate_var.get())
            fba        = float(self.mon_fba_var.get())
            min_profit = float(self.mon_min_profit_var.get())
        except ValueError:
            messagebox.showwarning("入力エラー", "数値を正しく入力してください。")
            return

        iid = self.mon_tree.insert(
            "", "end",
            values=(asin, f"¥{purchase:,.0f}", "—", "—", "—", "未確認"),
            tags=("loading",))
        self._monitor_items[iid] = {
            "asin": asin, "purchase": purchase,
            "ref_rate": ref_rate, "fba": fba, "min_profit": min_profit,
        }
        self.mon_asin_var.set("")
        self.mon_purchase_var.set("")

    def _delete_selected(self):
        for iid in self.mon_tree.selection():
            self.mon_tree.delete(iid)
            self._monitor_items.pop(iid, None)

    def _check_all(self):
        if not self._monitor_items:
            messagebox.showinfo("情報", "監視リストが空です。商品を追加してください。")
            return
        if not self.api_key:
            messagebox.showerror("APIキーエラー",
                "KeepaのAPIキーを入力して保存してください。")
            return
        self.check_all_btn.config(state="disabled", text="確認中...")
        self.mon_status_var.set("Keepaに問い合わせ中...")

        items = dict(self._monitor_items)

        def worker():
            results = {}
            for iid, info in items.items():
                try:
                    price, title = fetch_keepa_price(self.api_key, info["asin"])
                    results[iid] = {"price": price, "title": title, "error": None}
                except Exception as e:
                    results[iid] = {"price": None, "title": None, "error": str(e)}
            self.root.after(0, self._on_check_done, results)

        threading.Thread(target=worker, daemon=True).start()

    def _on_check_done(self, results):
        self.check_all_btn.config(state="normal", text="  全件 最安値チェック  ")
        updated = 0
        for iid, res in results.items():
            info = self._monitor_items.get(iid)
            if info is None:
                continue

            purchase   = info["purchase"]
            ref_rate   = info["ref_rate"]
            fba        = info["fba"]
            min_profit = info["min_profit"]

            if res["error"]:
                self.mon_tree.item(iid, values=(
                    info["asin"], f"¥{purchase:,.0f}",
                    "エラー", "—", "—", res["error"][:15]), tags=("ng",))
                continue

            market_price = res["price"]
            if market_price is None:
                self.mon_tree.item(iid, values=(
                    info["asin"], f"¥{purchase:,.0f}",
                    "在庫なし", "—", "—", "確認不可"), tags=("warning",))
                continue

            # 最低販売価格（利益確保ライン）
            min_selling = calc_min_selling_price(purchase, ref_rate, fba, 0, min_profit)
            if min_selling is None:
                tag, verdict = "ng", "設定エラー"
                rec_price_str = "—"
                rate_str = "—"
            else:
                # 推奨価格 = 競合最安値-1円 or 最低価格の高い方
                rec_price = max(market_price - 1, min_selling)
                rec_price = round(rec_price)

                _, profit, profit_rate = calc(purchase, rec_price, ref_rate, fba, 0)

                rec_price_str = f"¥{rec_price:,}"
                rate_str = f"{profit_rate:.1f}%"

                if rec_price <= market_price - 1:
                    tag = "ok"
                    verdict = "値下げ可"
                elif rec_price == market_price:
                    tag = "ok"
                    verdict = "同価格"
                else:
                    tag = "warning"
                    verdict = f"最低¥{min_selling:,.0f}"

            self.mon_tree.item(iid, values=(
                info["asin"],
                f"¥{purchase:,.0f}",
                f"¥{market_price:,.0f}",
                rec_price_str,
                rate_str,
                verdict,
            ), tags=(tag,))
            updated += 1

        now = datetime.datetime.now().strftime("%H:%M:%S")
        self.mon_status_var.set(f"最終確認: {now}  ({updated}件更新)")

    # ── Keepa 連携（利益計算タブ用） ──────────────────────────────────────────

    def _fetch_price(self):
        if not self.api_key:
            messagebox.showerror("APIキーエラー",
                f"keepa_config.json が見つかりません。\n"
                "api_key を設定してください。")
            return

        asin = extract_asin(self.asin_input.get())
        if not asin:
            messagebox.showwarning("入力エラー",
                "ASINまたはAmazonのURLを入力してください。\n"
                "例: B08XYZ1234  または  https://www.amazon.co.jp/dp/B08XYZ1234")
            return

        self.fetch_btn.config(state="disabled", text="取得中...")
        self.keepa_status.set("Keepaに問い合わせ中...")
        self.product_title.set("")

        def worker():
            try:
                price, title = fetch_keepa_price(self.api_key, asin)
                self.root.after(0, self._on_price_fetched, price, title)
            except Exception as e:
                self.root.after(0, self._on_fetch_error, str(e))

        threading.Thread(target=worker, daemon=True).start()

    def _on_price_fetched(self, price, title):
        self.fetch_btn.config(state="normal", text="価格取得")
        if price is None:
            self.keepa_status.set("価格が取得できませんでした（在庫なし等）")
            return
        self.selling.set(str(int(price)))
        self.keepa_status.set(f"取得成功: ¥{int(price):,}")
        if title:
            short = title[:50] + ("…" if len(title) > 50 else "")
            self.product_title.set(short)

    def _on_fetch_error(self, err):
        self.fetch_btn.config(state="normal", text="価格取得")
        self.keepa_status.set(f"エラー: {err}")

    # ── ロジック ─────────────────────────────────────────────────────────────

    def _on_platform(self):
        d = PLATFORM_DEFAULTS[self.platform.get()]
        self.referral_rate.set(str(d["referral_rate"]))
        self.fba_fee.set(str(d["fba_fee"]))
        self.shipping.set(str(d["shipping"]))
        self._refresh()

    def _float(self, var, default=0.0):
        try:
            return float(var.get().replace(",", "").strip())
        except ValueError:
            return default

    def _refresh(self, *_):
        purchase = self._float(self.purchase)
        selling  = self._float(self.selling)
        ref_rate = self._float(self.referral_rate, 15.0)
        fba      = self._float(self.fba_fee)
        ship     = self._float(self.shipping)

        referral_fee, profit, profit_rate = calc(
            purchase, selling, ref_rate, fba, ship)

        self.r_referral.set(f"¥ {referral_fee:,.0f}")
        self.r_fba.set(f"¥ {fba:,.0f}")
        self.r_shipping.set(f"¥ {ship:,.0f}")
        self.r_profit.set(f"¥ {profit:,.0f}")
        self.r_rate.set(f"{profit_rate:.1f} %")

        if purchase > 0 and selling > 0:
            color = "#27AE60" if profit > 0 else "#E74C3C"
            self.r_status.set("▲ 黒字" if profit > 0 else "▼ 赤字")
            self.profit_lbl.config(fg=color)
            self.status_lbl.config(fg=color)
        else:
            self.r_status.set("")

    def _clear(self):
        self.purchase.set("")
        self.selling.set("")
        self.asin_input.set("")
        self.product_title.set("")
        self.keepa_status.set("")
        self._on_platform()

    def _save(self):
        if not EXCEL_AVAILABLE:
            messagebox.showerror("エラー",
                "openpyxl がインストールされていません。\n"
                "コマンドプロンプトで:\n  pip install openpyxl\nを実行してください。")
            return

        purchase = self._float(self.purchase)
        selling  = self._float(self.selling)
        if purchase == 0 or selling == 0:
            messagebox.showwarning("入力エラー", "仕入れ値と販売価格を入力してください。")
            return

        ref_rate = self._float(self.referral_rate, 15.0)
        fba      = self._float(self.fba_fee)
        ship     = self._float(self.shipping)
        referral_fee, profit, profit_rate = calc(purchase, selling, ref_rate, fba, ship)

        if os.path.exists(EXCEL_FILE):
            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "利益計算"
            headers = ["日付", "プラットフォーム", "ASIN", "商品名",
                       "仕入れ値", "販売価格", "手数料率(%)",
                       "カテゴリー手数料", "FBA配送料", "送料", "利益", "利益率(%)"]
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2C3E50")
                cell.alignment = Alignment(horizontal="center")
            for i, w in enumerate([12, 16, 12, 30, 12, 12, 11, 14, 12, 10, 12, 11], 1):
                ws.column_dimensions[
                    openpyxl.utils.get_column_letter(i)].width = w

        asin_val = extract_asin(self.asin_input.get()) or ""
        title_val = self.product_title.get()

        row_data = [
            datetime.datetime.now().strftime("%Y/%m/%d"),
            self.platform.get(),
            asin_val, title_val,
            purchase, selling, ref_rate,
            round(referral_fee), round(fba), round(ship),
            round(profit), round(profit_rate, 1),
        ]
        ws.append(row_data)

        last_row = ws.max_row
        profit_cell = ws.cell(row=last_row, column=11)
        if profit > 0:
            profit_cell.fill = PatternFill("solid", fgColor="D5F5E3")
            profit_cell.font = Font(color="1E8449", bold=True)
        else:
            profit_cell.fill = PatternFill("solid", fgColor="FADBD8")
            profit_cell.font = Font(color="C0392B", bold=True)

        wb.save(EXCEL_FILE)
        messagebox.showinfo("保存完了",
            f"Excelに追記しました！\nファイル: {os.path.abspath(EXCEL_FILE)}")


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.mainloop()
